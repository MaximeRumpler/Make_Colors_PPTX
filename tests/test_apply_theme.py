import json
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

import apply_theme
import build_theme
import preview_deck
from _colour import load_theme
from conftest import make_clean_dir

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _deck(path, fixed_colour=False):
    """A small existing deck, with a shape named like a Make_Style role tag."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(3), Inches(1))
    box.name = "card#1 > Title"
    if fixed_colour:
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xC0, 0x30, 0x30)
    prs.save(path)
    return path


@pytest.fixture
def themes(tmp_path):
    roles = {"name": "Demo Harbour", "primary": "#1D4E89", "secondary": "#E07A1F", "others": ["#2A9D8F"],
             "fonts": {"major": "Georgia", "minor": "Arial"}}
    rep = build_theme.build(roles, tmp_path / "brand", thmx=True)
    return {v: rep["variants"][v] for v in ("light", "dark")}


def test_light_theme_on_an_existing_deck(tmp_path, themes):
    deck = _deck(tmp_path / "deck.pptx")
    before = deck.read_bytes()
    res = apply_theme.apply(deck, themes["light"]["files"][0])
    assert res["ok"] and not res["dark"] and res["fixed_colours"] == 0
    assert deck.read_bytes() == before  # the input is never touched
    out = load_theme(res["output"])
    assert out["colors"] == themes["light"]["colors"]
    assert out["fonts"] == {"major": "Georgia", "minor": "Arial"}
    with zipfile.ZipFile(deck) as a, zipfile.ZipFile(res["output"]) as b:
        assert a.namelist() == b.namelist()
        assert a.read("ppt/slides/slide1.xml") == b.read("ppt/slides/slide1.xml")  # shapes and role tags kept
    assert Presentation(res["output"]).slides[0].shapes[0].name == "card#1 > Title"


def test_dark_theme_swaps_the_colour_map(tmp_path, themes):
    res = apply_theme.apply(_deck(tmp_path / "deck.pptx"), themes["dark"]["files"][0], tmp_path / "dark.pptx")
    assert res["ok"] and res["dark"] and res["colour_map"] == "from the theme"
    with zipfile.ZipFile(res["output"]) as z:
        cm = etree.fromstring(z.read("ppt/slideMasters/slideMaster1.xml")).find(P + "clrMap")
    assert cm.get("bg1") == "dk1" and cm.get("tx1") == "lt1"


def test_thmx_keeps_the_deck_colour_map_and_counts_fixed_colours(tmp_path, themes):
    res = apply_theme.apply(_deck(tmp_path / "deck.pptx", fixed_colour=True), themes["light"]["files"][1])
    assert res["colour_map"] == "kept" and res["fixed_colours"] == 1
    assert any("fixed colours" in n for n in res["notes"])


def test_font_fallback_and_refusals(tmp_path, themes, capsys):
    roles = {"name": "Demo Font", "primary": "#1D4E89", "fonts": {"major": "Demo Display Sans"}, "variants": ["light"]}
    potx = build_theme.build(roles, tmp_path / "brand")["variants"]["light"]["files"][0]
    deck = _deck(tmp_path / "deck.pptx")
    res = apply_theme.apply(deck, potx, fallback="Arial")
    assert res["fonts"] == {"major": "Arial", "minor": "Arial"} and res["notes"]
    assert apply_theme.main([str(deck), "--theme", potx, "--out", str(deck)]) == 2
    assert "overwrite" in capsys.readouterr().out
    bad = tmp_path / "bad.pptx"
    bad.write_text("not a zip")
    assert apply_theme.main([str(bad), "--theme", potx]) == 2


def test_make_clean_deck_then_theme(tmp_path, themes):
    """Clean first, theme later (or after Make_Style): the same check as a build with the "theme" key."""
    mc = make_clean_dir()
    if not mc:
        pytest.skip("Make_Clean_PPTX not available")
    neutral = preview_deck.preview(themes["light"]["files"][0], out_dir=tmp_path / "built", make_clean=mc,
                                   render=False)
    assert "error" not in neutral, neutral
    res = preview_deck.preview(themes["light"]["files"][0], pptx=neutral["decks"]["before"],
                               out_dir=tmp_path / "prev", make_clean=mc, render=False)
    assert "error" not in res, res
    assert res["theme_report"]["fixed_colours"] == 0  # Make_Clean decks use theme colours only
    assert res["lint"].get("errors") == []
    contrast = res["lint"].get("contrast")
    if contrast is not None:
        assert contrast["failing_shapes"] == 0 and contrast["theme_findings"] == []
    assert json.dumps(res)
